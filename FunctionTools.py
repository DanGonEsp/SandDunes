from paraview.simple import *
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
import matplotlib.pyplot as plt
import matplotlib as mpl
import glob
import os
import csv
import sys


# ============================================================
# FUNCTION
# ============================================================

def calculate_contour(folder, dim, data_name):

    # ========================================================
    # CHECK INPUTS
    # ========================================================

    if dim not in [2, 3]:
        raise ValueError("ERROR: dim must be either 2 or 3. Code stopped.")

    if not os.path.isdir(folder):
        raise ValueError(f"ERROR: Folder does not exist: {folder}")

    if not data_name:
        raise ValueError("ERROR: Data folder name cannot be empty.")

    
    # ========================================================
    # OUTPUT FOLDER
    # ========================================================

    output_folder = os.path.join(folder, data_name)

    os.makedirs(output_folder, exist_ok=True)

    print("Input folder :", folder)
    print("Output folder:", output_folder)
    print("Dimension    :", dim)


    # ========================================================
    # FIND FILES
    # ========================================================

    pvtu_pattern = os.path.join(folder, "*.pvtu")
    vtu_pattern = os.path.join(folder, "*.vtu")

    pvtu_files = sorted(glob.glob(pvtu_pattern))
    vtu_files = sorted(glob.glob(vtu_pattern))

    # ========================================================
    # SELECT FILE TYPE AND READER
    # ========================================================

    if len(pvtu_files) > 0:

        files = pvtu_files

        print("Detected partitioned VTU data.")
        print("Using .pvtu files.")
        print("Found", len(files), "files")

        solution = XMLPartitionedUnstructuredGridReader(registrationName="Solution", FileName=files)

    elif len(vtu_files) > 0:

        files = vtu_files

        print("Detected standard VTU data.")
        print("Using .vtu files.")
        print("Found", len(files), "files")

        solution = XMLUnstructuredGridReader(
            registrationName="Solution", FileName=files
        )

    else:

        raise RuntimeError("ERROR: No .pvtu or .vtu files found!")


    # ========================================================
    # OUTPUT FILE
    # ========================================================

    OUTPUT_FILE = os.path.join(
        output_folder,
        f"Contour{dim}D.csv"
    )


    # ========================================================
    # READ FILES
    # ========================================================

    solution.UpdatePipeline()

    # ========================================================
    # TIME STEPS
    # ========================================================

    animationScene = GetAnimationScene()
    animationScene.UpdateAnimationUsingDataTimeSteps()

    timesteps = animationScene.TimeKeeper.TimestepValues

    print("Number of timesteps:", len(timesteps))


    # ========================================================
    # CREATE CONTOUR
    # ========================================================

    contour1 = Contour(Input=solution)

    contour1.ContourBy = ["POINTS", "c"]

    contour1.Isosurfaces = [0.5]

    contour1.UpdatePipeline()


    # ========================================================
    # OPEN CSV
    # ========================================================

    if os.path.exists(OUTPUT_FILE):
        print("WARNING: Output file already exists and will be overwritten:")
        print(OUTPUT_FILE)

    with open(OUTPUT_FILE, "w", newline="") as csvfile:

        writer = csv.writer(csvfile)

        if dim == 2:
            writer.writerow([
                "TimeStep", "Time", "X", "Y", "Z","c", "p", "u", "v","velocity_x", "velocity_y", "velocity_z"])

        elif dim == 3:
            writer.writerow(["TimeStep", "Time", "X", "Y", "Z","c", "p", "u", "v", "w","velocity_x", "velocity_y", "velocity_z"])


        # ====================================================
        # LOOP OVER TIME
        # ====================================================

        data_found = False
        
        for i, time in enumerate(timesteps):

            #print("Processing timestep", i, "/", len(timesteps) - 1)

            solution.UpdatePipeline(time)
            contour1.UpdatePipeline(time)

            data_contour = servermanager.Fetch(contour1)

            if data_contour is None:
                print("WARNING: No data")
                continue

            number_of_points = data_contour.GetNumberOfPoints()

            #print("Contour points:", number_of_points)

            if number_of_points == 0:
                print("WARNING: Empty contour at timestep", i)
                continue
                
            data_found = True

            # =================================================
            # DATA ARRAYS
            # =================================================

            point_data = data_contour.GetPointData()

            c_array = point_data.GetArray("c")
            p_array = point_data.GetArray("p")
            u_array = point_data.GetArray("u")
            v_array = point_data.GetArray("v")

            if dim == 3:
                w_array = point_data.GetArray("w")

            velocity_array = point_data.GetArray("velocity")


            # =================================================
            # WRITE POINTS
            # =================================================

            for point_id in range(number_of_points):

                x, y, z = data_contour.GetPoint(point_id)


                # ---------------------------------------------
                # Scalars
                # ---------------------------------------------

                c = c_array.GetTuple1(point_id) if c_array is not None else ""

                p = p_array.GetTuple1(point_id) if p_array is not None else ""

                u = u_array.GetTuple1(point_id) if u_array is not None else ""

                v = v_array.GetTuple1(point_id) if v_array is not None else ""

                if dim == 3:
                    w = w_array.GetTuple1(point_id) if w_array is not None else ""


                # ---------------------------------------------
                # Velocity vector
                # ---------------------------------------------

                if velocity_array is not None:

                    velocity = velocity_array.GetTuple(point_id)

                    velocity_x = velocity[0]
                    velocity_y = velocity[1]

                    if len(velocity) > 2:
                        velocity_z = velocity[2]
                    else:
                        velocity_z = 0.0

                else:

                    velocity_x = ""
                    velocity_y = ""
                    velocity_z = ""


                # ---------------------------------------------
                # Write row
                # ---------------------------------------------

                if dim == 2:

                    writer.writerow([i, time, x, y, z, c, p, u, v,velocity_x, velocity_y, velocity_z])

                elif dim == 3:

                    writer.writerow([i, time, x, y, z, c, p, u, v, w,velocity_x, velocity_y, velocity_z])

	# ========================================================
    # CHECK DATA
    # ========================================================

    if not data_found:
        print("")
        print("WARNING: No contour data was found.")
        
    # ========================================================
    # CLEANUP
    # ========================================================

    Delete(contour1)
    Delete(solution)


    # ========================================================
    # FINISHED
    # ========================================================

    print("")
    print("============================================")
    if data_found:
        print("DONE")
    else:
        print("DONE - NO CONTOUR DATA")
    print("============================================")
    print("Files found :", len(files))
    print("Output file :", OUTPUT_FILE)
    print("============================================")

    if data_found:
        return OUTPUT_FILE
    else:
        return None
# ============================================================
# PLOT CONTOUR
# ============================================================

def plot_contour(folder, dim, data_name, num_contours):

    # ========================================================
    # CHECK INPUTS
    # ========================================================

    if dim != 2: raise ValueError("ERROR: plot_contour is currently implemented only for dim = 2.")
    if not os.path.isdir(folder): raise ValueError(f"ERROR: Folder does not exist: {folder}")
    if not data_name: raise ValueError("ERROR: Data folder name cannot be empty.")
    if num_contours <= 0: raise ValueError("ERROR: num_contours must be greater than zero.")

    # ========================================================
    # INPUT / OUTPUT FILES
    # ========================================================

    data_folder = os.path.join(folder, data_name)
    input_file = os.path.join(data_folder, f"Contour{dim}D.csv")
    output_file = os.path.join(data_folder, f"ContourEvolution{dim}D.png")

    if not os.path.isfile(input_file):
        print(f"Contour file does not exist: {input_file}")
        print("Creating contour file...")
        calculate_contour(folder, dim, data_name)
        if not os.path.isfile(input_file):
            raise RuntimeError(f"ERROR: Failed to create contour file: {input_file}")

    print("")
    print("============================================")
    print("PLOT CONTOUR")
    print("============================================")
    print("Input file      :", input_file)
    print("Output file     :", output_file)
    print("Requested plots :", num_contours)

    # ========================================================
    # READ CONTOUR DATA
    # ========================================================

    contours = {}

    with open(input_file, "r", newline="") as csvfile:
        reader = csv.DictReader(csvfile)
        required_columns = ["TimeStep", "Time", "X", "Y"]

        if reader.fieldnames is None: raise RuntimeError(f"ERROR: Empty CSV file: {input_file}")

        for column in required_columns:
            if column not in reader.fieldnames: raise RuntimeError(f"ERROR: Column '{column}' was not found in: {input_file}")

        for row in reader:
            timestep = int(row["TimeStep"])
            time = float(row["Time"])
            x = float(row["X"])
            y = float(row["Y"])
            if timestep not in contours: contours[timestep] = {"time": time, "x": [], "y": []}
            contours[timestep]["x"].append(x)
            contours[timestep]["y"].append(y)

    # ========================================================
    # CHECK DATA
    # ========================================================

    if len(contours) == 0: raise RuntimeError("ERROR: No contour data was found.")

    timesteps = sorted(contours.keys())
    print("Available plots :", len(timesteps))

    # ========================================================
    # SELECT TIMESTEPS
    # ========================================================

    num_selected = min(num_contours, len(timesteps))

    if num_selected == 1:
        selected_timesteps = [timesteps[0]]
    else:
        selected_timesteps = []
        for i in range(num_selected):
            index = round(i * (len(timesteps) - 1) / (num_selected - 1))
            selected_timesteps.append(timesteps[index])

    print("Selected plots  :", len(selected_timesteps))

    # ========================================================
    # SELECTED TIMES
    # ========================================================

    selected_times = [contours[timestep]["time"] for timestep in selected_timesteps]
    time_min = min(selected_times)
    time_max = max(selected_times)

    # ========================================================
    # COLORMAP
    # ========================================================

    if time_min == time_max: norm = mpl.colors.Normalize(vmin=time_min - 0.5, vmax=time_max + 0.5)
    else: norm = mpl.colors.Normalize(vmin=time_min, vmax=time_max)

    cmap = plt.get_cmap("viridis")

    # ========================================================
    # CREATE FIGURE
    # ========================================================

    fig, ax = plt.subplots(figsize=(10, 5))

    # ========================================================
    # PLOT CONTOURS
    # ========================================================

    for timestep in selected_timesteps:
        time = contours[timestep]["time"]
        points = sorted(zip(contours[timestep]["x"], contours[timestep]["y"]), key=lambda point: point[0])
        x_sorted = [point[0] for point in points]
        y_sorted = [point[1] for point in points]
        ax.plot(x_sorted, y_sorted, color=cmap(norm(time)), linewidth=1.2)

    # ========================================================
    # COLORBAR
    # ========================================================

    scalar_map = mpl.cm.ScalarMappable(norm=norm, cmap=cmap)
    scalar_map.set_array([])
    colorbar = fig.colorbar(scalar_map, ax=ax)
    colorbar.set_label("Time")

    # ========================================================
    # AXES
    # ========================================================

    ax.set_xlabel("X")
    ax.set_ylabel("Y")
    ax.set_aspect("equal")

    # ========================================================
    # SAVE FIGURE
    # ========================================================

    fig.tight_layout()
    fig.savefig(output_file, dpi=300, bbox_inches="tight")
    plt.close(fig)

    # ========================================================
    # FINISHED
    # ========================================================

    print("")
    print("============================================")
    print("DONE")
    print("============================================")
    print("Contours plotted:", len(selected_timesteps))
    print("Initial time    :", time_min)
    print("Final time      :", time_max)
    print("Output file     :", output_file)
    print("============================================")

    return output_file

# ============================================================
# SCALABILITY
# ============================================================

def StrongScalability(folder, step, level):

    # ========================================================
    # CHECK INPUTS
    # ========================================================

    if not os.path.isdir(folder):
        raise ValueError(f"ERROR: Folder does not exist: {folder}")

    # ========================================================
    # FIND PARALLEL CASES
    # ========================================================

    parallel_pattern = os.path.join(folder,f"*lev{level}*")

    parallel_folders = sorted(glob.glob(parallel_pattern))

    if len(parallel_folders) == 0:
        raise RuntimeError("ERROR: No parallel simulation folders were found!")

    # ========================================================
    # CHECK FILES
    # ========================================================

    cases = []

    for parallel_folder in parallel_folders:

        if not os.path.isdir(parallel_folder):
            continue

        folder_name = os.path.basename(parallel_folder)

        ranks = int(folder_name.split("_")[-1])

        integral_file = os.path.join(parallel_folder,"Integral.txt")

        if not os.path.isfile(integral_file):
            raise RuntimeError(f"ERROR: Integral.txt does not exist: {integral_file}")

        # ====================================================
        # READ INTEGRAL
        # ====================================================

        numstep = None
        total_work_time = None
        limex_total = None
        limex_success = None
        limex_fail = None
        linear_calls = None
        linear_steps = None

        with open(integral_file, "r") as file:

            for line in file:

                if line.startswith("Step") or line.startswith("-"):
                    continue

                columns = line.split()

                if len(columns) < 11:
                    continue

                current_step = int(columns[0])

                if current_step == step:

                    numstep = current_step
                    total_work_time = float(columns[5])
                    limex_total = int(columns[6])
                    limex_success = int(columns[7])
                    limex_fail = int(columns[8])
                    linear_calls = int(columns[9])
                    linear_steps = int(columns[10])

        if total_work_time is None:
            raise RuntimeError(f"ERROR: Step {step} was not found in: {integral_file}")
        if limex_total != limex_success + limex_fail:
            raise RuntimeError(f"ERROR: LIMEX counters are inconsistent in: {integral_file}")
        if limex_total == 0:
            raise RuntimeError(f"ERROR: LIMEX total steps is zero in: {integral_file}")
            
        limex_time = total_work_time / limex_total

        print("Ranks:", ranks," Step:", numstep," Total work time:", total_work_time)

        cases.append([ranks, numstep, total_work_time, limex_total, limex_success, limex_fail, limex_time, linear_calls, linear_steps, parallel_folder, integral_file])
    cases.sort(key=lambda case: case[0])
    reference_time = cases[0][2]
    for case in cases:
        if case[7] == 0:
            raise RuntimeError(f"ERROR: Linear solver calls is zero in: {case[10]}")
        speedup = reference_time / case[2]
        avg_linear_steps = case[8] / case[7]

        case.append(speedup)
        case.append(avg_linear_steps)
        
    print("")
    print("==========================================================================================================")
    print("STRONG SCALABILITY RESULTS")
    print("==========================================================================================================")
    print(f"{'PE':>6} {'Ttotal(s)':>14} {'NTimeSteps':>14} {'tTimeStep(s)':>14} {'LinCalls':>12} {'LinSteps':>12} {'AvgLinSteps':>14} {'Speedup':>10}")
    print("----------------------------------------------------------------------------------------------------------")
    for case in cases:
        ntime_steps = f"{case[4]} ({case[5]})"
        print(f"{case[0]:6d} {case[2]:14.6f} {ntime_steps:>12} {case[6]:14.6f} {case[7]:12d} {case[8]:12d} {case[12]:14.3f} {case[11]:10.3f}")
    print("==========================================================================================================")
    
	# ========================================================
    # WRITE EXCEL FILE
    # ========================================================

    excel_file = os.path.join(folder,"StrongScalability.xlsx")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Strong Scalability"

    headers = ["PE", "Ttotal (s)", "NTimeSteps", "tTimeStep (s)", "LinCalls", "LinSteps", "AvgLinSteps", "Speedup"]
    worksheet.append(headers)

    for case in cases:
        ntime_steps = f"{case[4]} ({case[5]})"
        worksheet.append([case[0], case[2], ntime_steps, case[6], case[7], case[8], case[12], case[11]])

    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(horizontal="center")

    worksheet.column_dimensions["A"].width = 10
    worksheet.column_dimensions["B"].width = 15
    worksheet.column_dimensions["C"].width = 15
    worksheet.column_dimensions["D"].width = 15
    worksheet.column_dimensions["E"].width = 15
    worksheet.column_dimensions["F"].width = 15
    worksheet.column_dimensions["G"].width = 15
    worksheet.column_dimensions["H"].width = 12

    for cell in worksheet["B"][1:]:
        cell.number_format = "0.000000"

    for cell in worksheet["D"][1:]:
        cell.number_format = "0.000000"

    for cell in worksheet["G"][1:]:
        cell.number_format = "0.000"

    for cell in worksheet["H"][1:]:
        cell.number_format = "0.000"

    workbook.save(excel_file)

    print("")
    print("Excel file:", excel_file)
        
    # ========================================================
    # FINISHED
    # ========================================================

    print("")

    print("============================================")
    print("SCALABILITY FILE CHECK")
    print("============================================")
    print("Cases found:", len(cases))
    print("Level:", level)
    print("Step:", step)
    print("All Integral.txt files exist.")
    print("============================================")

    return cases
# ============================================================
# WEAK SCALABILITY
# ============================================================

def WeakScalability(folder, step, factor):

    # ========================================================
    # CHECK INPUTS
    # ========================================================

    if not os.path.isdir(folder):
        raise ValueError(f"ERROR: Folder does not exist: {folder}")

    if factor <= 0:
        raise ValueError("ERROR: factor must be greater than zero.")

    # ========================================================
    # FIND PARALLEL CASES
    # ========================================================

    parallel_pattern = os.path.join(folder,"*lev*")

    parallel_folders = sorted(glob.glob(parallel_pattern))

    if len(parallel_folders) == 0:
        raise RuntimeError("ERROR: No parallel simulation folders were found!")



    # ========================================================
    # FIND INITIAL LEVEL AND n0
    # ========================================================

    lev0 = None
    n0 = None

    for parallel_folder in parallel_folders:

        if not os.path.isdir(parallel_folder):
            continue

        folder_name = os.path.basename(parallel_folder)

        lev = int(folder_name.split("lev")[1].split("_")[0])

        parallel_rank = int(folder_name.split("_")[-1])

        if lev0 is None or lev < lev0 or (lev == lev0 and parallel_rank < n0):
            lev0 = lev
            n0 = parallel_rank

    if lev0 is None:
        raise RuntimeError("ERROR: No valid weak scalability levels were found!")

    # ========================================================
    # CHECK FILES
    # ========================================================

    cases = []

    for parallel_folder in parallel_folders:

        if not os.path.isdir(parallel_folder):
            continue

        folder_name = os.path.basename(parallel_folder)

        lev = int(folder_name.split("lev")[1].split("_")[0])

        parallel_rank = int(folder_name.split("_")[-1])

        expected_rank = n0 * factor**(lev - lev0)

        if parallel_rank != expected_rank:
            continue

        integral_file = os.path.join(parallel_folder,"Integral.txt")

        if not os.path.isfile(integral_file):
            raise RuntimeError(f"ERROR: Integral.txt does not exist: {integral_file}")

        # ====================================================
        # READ INTEGRAL
        # ====================================================

        numstep = None
        total_work_time = None
        time_total = None
        time_success = None
        time_fail = None
        linear_calls = None
        linear_steps = None

        with open(integral_file, "r") as file:

            for line in file:

                if line.startswith("Step") or line.startswith("-"):
                    continue

                columns = line.split()

                if len(columns) < 11:
                    continue

                current_step = int(columns[0])

                if current_step == step:

                    numstep = current_step
                    total_work_time = float(columns[5])
                    time_total = int(columns[6])
                    time_success = int(columns[7])
                    time_fail = int(columns[8])
                    linear_calls = int(columns[9])
                    linear_steps = int(columns[10])

        if total_work_time is None:
            raise RuntimeError(f"ERROR: Step {step} was not found in: {integral_file}")
        if time_total != time_success + time_fail:
            raise RuntimeError(f"ERROR: Time-step counters are inconsistent in: {integral_file}")
        if time_total == 0:
            raise RuntimeError(f"ERROR: Total number of time steps is zero in: {integral_file}")
        if linear_calls == 0:
            raise RuntimeError(f"ERROR: Linear solver calls is zero in: {integral_file}")
            
        time_step_time = total_work_time / time_total
        avg_linear_steps = linear_steps / linear_calls

        print("Lev:", lev," Parallel ranks:", parallel_rank," Step:", numstep," Total work time:", total_work_time)

        cases.append([lev, parallel_rank, numstep, total_work_time, time_total, time_success, time_fail, time_step_time, linear_calls, linear_steps, avg_linear_steps, parallel_folder, integral_file])
        
    # ========================================================
    # SORT AND CALCULATE WEAK EFFICIENCY
    # ========================================================

    if len(cases) == 0:
        raise RuntimeError("ERROR: No valid weak scalability cases were found!")

    cases.sort(key=lambda case: case[0])

    reference_time = cases[0][3]

    for case in cases:
        raw_efficiency = reference_time / case[3]
        cfl_factor = 2**(case[0] - lev0)
        cfl_time = case[3] / cfl_factor
        cfl_efficiency = reference_time / cfl_time

        case.append(raw_efficiency)
        case.append(cfl_time)
        case.append(cfl_efficiency)
        
    # ========================================================
    # PRINT RESULTS
    # ========================================================

    print("")
    print("================================================================================================================")
    print("WEAK SCALABILITY RESULTS")
    print("================================================================================================================")
    print(f"{'Level':>7} {'PE':>6} {'Ttotal(s)':>14} {'NTimeSteps':>14} {'tTimeStep(s)':>14} {'LinCalls':>12} {'LinSteps':>12} {'AvgLinSteps':>14} {'RawEff':>10} {'CFLTime(s)':>12} {'CFLEff':>10}")
    print("----------------------------------------------------------------------------------------------------------------")

    for case in cases:
        ntime_steps = f"{case[5]} ({case[6]})"
        print(f"{case[0]:7d} {case[1]:6d} {case[3]:14.6f} {ntime_steps:>14} {case[7]:14.6f} {case[8]:12d} {case[9]:12d} {case[10]:14.3f} {case[13]:10.3f} {case[14]:12.6f} {case[15]:10.3f}")

    print("================================================================================================================")
    # ========================================================
    # WRITE EXCEL FILE
    # ========================================================

    excel_file = os.path.join(folder,"WeakScalability.xlsx")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Weak Scalability"

    headers = ["Level", "PE", "Ttotal (s)", "NTimeSteps", "tTimeStep (s)", "LinCalls", "LinSteps", "AvgLinSteps", "RawEfficiency", "CFLTime (s)", "CFLEfficiency"]
    worksheet.append(headers)

    for case in cases:
        ntime_steps = f"{case[5]} ({case[6]})"
        worksheet.append([case[0], case[1], case[3], ntime_steps, case[7], case[8], case[9], case[10], case[13], case[14], case[15]])

    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(horizontal="center")

    worksheet.column_dimensions["A"].width = 10
    worksheet.column_dimensions["B"].width = 10
    worksheet.column_dimensions["C"].width = 15
    worksheet.column_dimensions["D"].width = 15
    worksheet.column_dimensions["E"].width = 15
    worksheet.column_dimensions["F"].width = 15
    worksheet.column_dimensions["G"].width = 15
    worksheet.column_dimensions["H"].width = 15
    worksheet.column_dimensions["I"].width = 12
    worksheet.column_dimensions["J"].width = 15
	worksheet.column_dimensions["K"].width = 15

    for cell in worksheet["C"][1:]:
        cell.number_format = "0.000000"

    for cell in worksheet["E"][1:]:
        cell.number_format = "0.000000"

    for cell in worksheet["H"][1:]:
        cell.number_format = "0.000"

    for cell in worksheet["I"][1:]:
        cell.number_format = "0.000"
        
    for cell in worksheet["I"][1:]:
        cell.number_format = "0.000"

    for cell in worksheet["J"][1:]:
        cell.number_format = "0.000000"

    for cell in worksheet["K"][1:]:
        cell.number_format = "0.000"

    workbook.save(excel_file)

    print("")
    print("Excel file:", excel_file)
    # ========================================================
    # FINISHED
    # ========================================================

    print("")

    print("============================================")
    print("WEAK SCALABILITY FILE CHECK")
    print("============================================")
    print("Cases found:", len(cases))
    print("Step:", step)
    print("Factor:", factor)
    print("Initial level:", lev0)
    print("Initial ranks:", n0)
    print("============================================")

    return cases


# ============================================================
# COMMAND-LINE INTERFACE
# ============================================================

if __name__ == "__main__":

    if len(sys.argv) < 2:
        raise RuntimeError("Usage: pvpython FunctionTools.py <function> <arguments>")

    function = sys.argv[1]

    if function == "contour":

        if len(sys.argv) != 5:
            raise RuntimeError("Usage: pvpython FunctionTools.py contour <folder> <dim> <data_name>")

        folder = sys.argv[2]
        dim = int(sys.argv[3])
        data_name = sys.argv[4]

        calculate_contour(folder, dim, data_name)
        
    elif function == "plot_contour":
        if len(sys.argv) != 6:
            raise RuntimeError("Usage: pvpython FunctionTools.py plot_contour <folder> <dim> <data_name> <num_contours>")
            
        folder = sys.argv[2]
        dim = int(sys.argv[3])
        data_name = sys.argv[4]
        num_contours = int(sys.argv[5])
        plot_contour(folder, dim, data_name, num_contours)

    elif function == "strong_scalability":

        if len(sys.argv) != 5:
                    raise RuntimeError("Usage: pvpython FunctionTools.py strong_scalability <folder> <step> <level>")

        folder = sys.argv[2]
        step = int(sys.argv[3])
        level = int(sys.argv[4])

        StrongScalability(folder, step, level)
    elif function == "weak_scalability":
        if len(sys.argv) != 5:
            raise RuntimeError("Usage: pvpython FunctionTools.py weak_scalability <folder> <step> <factor>")

        folder = sys.argv[2]
        step = int(sys.argv[3])
        factor = int(sys.argv[4])
        
        WeakScalability(folder, step, factor)

    else:
        raise RuntimeError(f"ERROR: Unknown function: {function}")

